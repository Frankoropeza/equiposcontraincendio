from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

ROJO = "C62828"; GRIS = "F2F3F5"
thin = Side(style="thin", color="BFC5CC")
BORDE = Border(left=thin, right=thin, top=thin, bottom=thin)
H = Font(bold=True, color="FFFFFF", size=10)
FILL = PatternFill("solid", fgColor=ROJO)
SUB = Font(bold=True, size=10)
FILL_SUB = PatternFill("solid", fgColor=GRIS)


def print_setup(ws, fila_encabezado, ncols, filas_totales):
    """Impresión: horizontal, ajustado al ancho de una hoja y con la cabecera
    repetida. Sin esto, la bitácora sale en 8 páginas y es inservible impresa."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{fila_encabezado}:{fila_encabezado}"
    ws.print_area = f"A1:{get_column_letter(ncols)}{filas_totales}"
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = ws.page_margins.bottom = 0.4

def encabezado(ws, titulo, subtitulo, ancho):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ancho)
    c = ws.cell(row=1, column=1, value=titulo)
    c.font = Font(bold=True, size=14, color=ROJO)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ancho)
    c2 = ws.cell(row=2, column=1, value=subtitulo)
    c2.font = Font(size=9, color="6B7280")
    campos = [("Razón social:", ""), ("Domicilio del centro de trabajo:", ""),
              ("Grado de riesgo de incendio:", "Ordinario / Alto"), ("Año:", ""),
              ("Responsable del registro:", "")]
    r = 4
    for etiqueta, val in campos:
        # La etiqueta ocupa A:B — la columna A es estrecha (es "No." de la tabla)
        # y sola cortaba el texto ("Razón s", "Domicil") al imprimir.
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        c_lab = ws.cell(row=r, column=1, value=etiqueta)
        c_lab.font = SUB
        c_lab.alignment = Alignment(horizontal="left", vertical="center")
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=min(ancho, 8))
        cv = ws.cell(row=r, column=3, value=val)
        cv.fill = FILL_SUB
        cv.border = BORDE
        r += 1
    return r + 1

def cabecera_tabla(ws, fila, columnas):
    for i, (titulo, ancho) in enumerate(columnas, start=1):
        c = ws.cell(row=fila, column=i, value=titulo)
        c.font = H; c.fill = FILL; c.border = BORDE
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.row_dimensions[fila].height = 30

def filas_vacias(ws, desde, n, ncols):
    for r in range(desde, desde + n):
        for c in range(1, ncols + 1):
            ws.cell(row=r, column=c).border = BORDE

def pie(ws, fila, ncols, textos):
    for i, t in enumerate(textos):
        ws.merge_cells(start_row=fila + i, start_column=1, end_row=fila + i, end_column=ncols)
        c = ws.cell(row=fila + i, column=1, value=t)
        c.font = Font(size=9, color="374151")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[fila + i].height = 26

# ── 1) Bitácora de revisión mensual de extintores ───────────────────────────
wb = Workbook()
ws = wb.active; ws.title = "Bitácora"
MESES = ["ENE","FEB","MAR","ABR","MAY","JUN","JUL","AGO","SEP","OCT","NOV","DIC"]
cols = [("No.", 6), ("Ubicación / área", 28), ("Agente", 14), ("Capacidad", 11),
        ("Identificación o serie", 20), ("Última recarga", 14), ("Prueba hidrostática", 16)]
cols += [(m, 5) for m in MESES]
cols += [("Observaciones", 30)]
fila = encabezado(ws, "Bitácora de revisión mensual de extintores",
                  "Formato de apoyo. La NOM-002-STPS-2010 exige revisión mensual del equipo y mantenimiento al menos una vez por año. equiposcontraincendio.com",
                  len(cols))
cabecera_tabla(ws, fila, cols)
filas_vacias(ws, fila + 1, 30, len(cols))
ws.freeze_panes = ws.cell(row=fila + 1, column=1)
print_setup(ws, fila, len(cols), fila + 34)
pie(ws, fila + 32, len(cols), [
 "Cómo se usa: una fila por extintor. Cada mes se marca la casilla del mes cuando se verificó el equipo; si algo falla, se anota en Observaciones y se corrige.",
 "Qué se revisa cada mes: que esté en su lugar y accesible, señalizado, a no más de 1.50 m de altura; manómetro dentro del rango; pasador y sello íntegros; cilindro sin daños ni corrosión; manguera y boquilla completas; etiqueta de servicio vigente.",
 "Este formato es de apoyo y no sustituye el programa ni la documentación que exige la normatividad aplicable a cada centro de trabajo.",
])
ws2 = wb.create_sheet("Qué se revisa")
ws2.column_dimensions["A"].width = 6; ws2.column_dimensions["B"].width = 86
ws2.cell(row=1, column=1, value="Puntos de la revisión mensual").font = Font(bold=True, size=13, color=ROJO)
PUNTOS = [
 "Está en su lugar, a la vista y sin nada que estorbe el acceso.",
 "La parte más alta del extintor queda a 1.50 m o menos sobre el nivel del piso.",
 "La señalización que lo identifica es visible desde la circulación del área.",
 "El manómetro marca dentro del rango de operación.",
 "El pasador y el sello de seguridad están puestos e íntegros.",
 "El cilindro no tiene golpes, abolladuras ni corrosión.",
 "La manguera y la boquilla están completas, sin cuarteaduras ni obstrucciones.",
 "La etiqueta de servicio tiene fecha y el mantenimiento sigue vigente (menos de un año).",
 "Se conoce la fecha en que toca la prueba hidrostática del cilindro.",
 "El agente corresponde al riesgo del área (clase K en cocina, CO2 en cuarto eléctrico).",
]
for i, p in enumerate(PUNTOS, start=3):
    ws2.cell(row=i, column=1, value=i - 2).alignment = Alignment(horizontal="center")
    ws2.cell(row=i, column=2, value=p).alignment = Alignment(wrap_text=True)
    ws2.row_dimensions[i].height = 22
ws2.page_setup.fitToWidth = 1
ws2.page_setup.fitToHeight = 0
ws2.sheet_properties.pageSetUpPr.fitToPage = True
wb.save("bitacora-revision-mensual-extintores.xlsx")

# ── 2) Censo de brigada de emergencia ───────────────────────────────────────
wb = Workbook(); ws = wb.active; ws.title = "Censo"
cols = [("No.", 6), ("Nombre completo", 30), ("Puesto", 20), ("Área", 18), ("Turno", 10),
        ("Brigada asignada", 22), ("Rol (jefe / integrante)", 18), ("Teléfono", 14),
        ("Curso recibido", 26), ("Fecha del curso", 14), ("Constancia (DC-3)", 16)]
fila = encabezado(ws, "Censo de brigada de emergencia",
                  "Formato de apoyo. La NOM-002-STPS-2010 establece requisitos de integración y capacitación de brigadas. equiposcontraincendio.com",
                  len(cols))
cabecera_tabla(ws, fila, cols)
filas_vacias(ws, fila + 1, 25, len(cols))
ws.freeze_panes = ws.cell(row=fila + 1, column=1)
print_setup(ws, fila, len(cols), fila + 29)
pie(ws, fila + 27, len(cols), [
 "Brigadas típicas: prevención y combate de incendio, evacuación, primeros auxilios, y comunicación. Anota una fila por persona; si alguien participa en dos brigadas, usa dos filas.",
 "Conserva copia de la constancia de cada curso: es lo que acredita la capacitación ante una verificación.",
 "Este formato es de apoyo y no sustituye el programa de capacitación ni la documentación que exige la normatividad aplicable.",
])
wb.save("censo-brigada-emergencia.xlsx")
print("xlsx listos")
