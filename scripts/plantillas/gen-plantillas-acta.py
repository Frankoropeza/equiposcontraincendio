from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

ROJO = RGBColor(0xC6, 0x28, 0x28)
doc = Document()
for s in doc.sections:
    s.top_margin = Cm(1.8); s.bottom_margin = Cm(1.8)
    s.left_margin = Cm(2); s.right_margin = Cm(2)
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10)

t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("ACTA DE SIMULACRO DE EVACUACIÓN"); r.bold = True; r.font.size = Pt(16); r.font.color.rgb = ROJO
s = doc.add_paragraph(); s.alignment = WD_ALIGN_PARAGRAPH.CENTER
rs = s.add_run("Formato de apoyo · equiposcontraincendio.com"); rs.font.size = Pt(8); rs.font.color.rgb = RGBColor(0x6B,0x72,0x80)

def titulo(txt):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(4)
    run = p.add_run(txt); run.bold = True; run.font.size = Pt(11); run.font.color.rgb = ROJO

def campos(pares, cols=2):
    tb = doc.add_table(rows=0, cols=cols * 2); tb.style = "Table Grid"; tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i in range(0, len(pares), cols):
        row = tb.add_row().cells
        for j in range(cols):
            idx = i + j
            etiqueta, ancho = (pares[idx], None) if idx < len(pares) else ("", None)
            c_lab = row[j * 2]; c_val = row[j * 2 + 1]
            pr = c_lab.paragraphs[0].add_run(etiqueta); pr.bold = True; pr.font.size = Pt(9)
            c_val.text = ""
    return tb

titulo("1. Datos del centro de trabajo")
campos(["Razón social:", "Giro o actividad:", "Domicilio:", "Alcaldía o municipio:",
        "Superficie construida (m²):", "Grado de riesgo de incendio:",
        "Número de ocupantes habitual:", "Responsable del inmueble:"])

titulo("2. Datos del ejercicio")
campos(["Fecha:", "Hipótesis (incendio, sismo, fuga…):",
        "Hora de inicio:", "Tipo (con aviso / sin aviso):",
        "Hora de término:", "Alcance (parcial / total):",
        "Áreas participantes:", "Apoyo externo (si hubo):"])

titulo("3. Tiempos registrados")
campos(["Hora de activación de la alarma:", "Hora en que se completó la evacuación:",
        "Tiempo total de evacuación:", "Tiempo hasta el conteo en el punto de reunión:",
        "Personas evacuadas:", "Brigadistas participantes:"])

titulo("4. Puntos observados durante el ejercicio")
obs = [
 "La alarma se escuchó en todas las áreas ocupadas.",
 "Las rutas de evacuación estaban libres de obstáculos.",
 "La señalización de rutas y salidas fue visible y suficiente.",
 "Las salidas de emergencia abrieron sin candados ni seguros.",
 "Los brigadistas cumplieron su función asignada.",
 "Se apoyó a personas con discapacidad o movilidad reducida.",
 "Se realizó el conteo de personas en el punto de reunión.",
 "La comunicación entre brigadas fue efectiva.",
 "Se controló el reingreso al inmueble hasta la indicación oficial.",
]
tb = doc.add_table(rows=1, cols=4); tb.style = "Table Grid"
tb.autofit = False
for row_cells, widths in [(tb.rows[0].cells, [Cm(1), Cm(11.5), Cm(2.3), Cm(2.3)])]:
    for c, w in zip(row_cells, widths):
        c.width = w
hdr = tb.rows[0].cells
for i, h in enumerate(["#", "Punto observado", "Sí", "No"]):
    rr = hdr[i].paragraphs[0].add_run(h); rr.bold = True; rr.font.size = Pt(9)
for i, o in enumerate(obs, start=1):
    row = tb.add_row().cells
    for c, w in zip(row, [Cm(1), Cm(11.5), Cm(2.3), Cm(2.3)]):
        c.width = w
    row[0].text = str(i); row[1].text = o; row[2].text = ""; row[3].text = ""
    for c in row: 
        for p in c.paragraphs:
            for run in p.runs: run.font.size = Pt(9)

titulo("5. Desarrollo y observaciones")
# Renglones para escribir: subrayado por borde, no por guiones bajos —los
# guiones se desbordan a la línea siguiente al convertir a PDF.
for _ in range(4):
    p = doc.add_paragraph("")
    p.paragraph_format.space_after = Pt(14)
    pPr = p._p.get_or_add_pPr()
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1"); bottom.set(qn("w:color"), "BFC5CC")
    pbdr.append(bottom); pPr.append(pbdr)

titulo("6. Áreas de oportunidad y acciones correctivas")
tb = doc.add_table(rows=1, cols=4); tb.style = "Table Grid"
hdr = tb.rows[0].cells
for i, h in enumerate(["Hallazgo", "Acción correctiva", "Responsable", "Fecha compromiso"]):
    rr = hdr[i].paragraphs[0].add_run(h); rr.bold = True; rr.font.size = Pt(9)
for _ in range(5):
    row = tb.add_row().cells
    for c in row: c.text = ""

titulo("7. Firmas")
tb = doc.add_table(rows=2, cols=3); tb.style = "Table Grid"
for i, h in enumerate(["Responsable del inmueble", "Jefe de brigada", "Testigo"]):
    rr = tb.rows[0].cells[i].paragraphs[0].add_run(h); rr.bold = True; rr.font.size = Pt(9)
for c in tb.rows[1].cells:
    c.text = "\n\n"

nota = doc.add_paragraph()
rn = nota.add_run(
 "Formato de apoyo para documentar el ejercicio. No sustituye el programa interno de protección civil "
 "ni la documentación que exija la autoridad competente en cada caso. Descargado de equiposcontraincendio.com"
)
rn.font.size = Pt(8); rn.font.color.rgb = RGBColor(0x6B,0x72,0x80)

doc.save("acta-simulacro-evacuacion.docx")
print("docx listo")
